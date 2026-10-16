from pathlib import Path
import json,re,html,sys,collections
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
sys.stdout.reconfigure(encoding='utf-8')
R=Path(__file__).parent
OUT=R/'AutoHaina_Справочник_данных_автомобилей.docx'
def load(n): return json.loads((R/n).read_text(encoding='utf-8'))
def mapping(n):return dict(x.split('|',1) for x in (R/n).read_text(encoding='utf-8').splitlines() if '|' in x)
T=mapping('translations.txt'); F=mapping('filter_translations.txt'); U=load('union.json'); filters=load('filters_en.json')['data']
files=['gas_en','ev_en','phev_en','erev_en','mpv_en','flagship_en','tesla_en','tank_en','prado_en']
C={f:load(f+'.json') for f in files}; codes={f:'К'+str(i+1) for i,f in enumerate(files)}
slugs=['mercedes-benz-glc-2023-0md6fj8ctt','dongfeng-yi-pai-e-007-2026-i0f2vdtu78','byd-seal-06-2027-rlfqgphcgp','aito-aito-aito-m7-2022-u7rg4duded','buick-century-2023-qfsavlr2no','voyah-car-voyah-taishan-2026-1gvt4f0acu','tesla-model-3-2023-vn62mvjh5s','tank-500-new-energy-2023-qfcs22ib3c','toyota-prado-imported-2024-q2pors2cmk']
names=['Mercedes Benz GLC 260 L','Dongfeng eπ007 650 Ultra','BYD Seal 06 DM i 230KM','AITO M7','Buick Century 48V','Voyah Taishan Max','Tesla Model 3','Tank 500 Hi4 T','Toyota Prado 4.0']
d=Document(); sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)
for sn in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=d.styles[sn];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
normal=d.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.08
for sn,size in [('Title',27),('Heading 1',18),('Heading 2',13),('Heading 3',11)]:
 st=d.styles[sn];st.font.size=Pt(size);st.paragraph_format.space_before=Pt(14);st.paragraph_format.space_after=Pt(7)
 st.paragraph_format.keep_with_next=True
hp=sec.header.paragraphs[0];hp.text='AUTOHAINA   /   СПРАВОЧНИК ДАННЫХ';hp.style='Caption';hp.runs[0].font.size=Pt(8);hp.runs[0].font.color.rgb=RGBColor(0,0,0)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('Проверено 14 октября 2026   •   ').font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
d.core_properties.title='Данные автомобилей AutoHaina для вашего сайта';d.core_properties.subject='Фильтры, характеристики, справочники и структура карточки автомобиля'
d.core_properties.author='';d.core_properties.keywords='AutoHaina, автомобили, характеристики, фильтры'
def p(t,style=None):
 para=d.add_paragraph(t,style)
 if t.startswith('Раздел источника:'):para.paragraph_format.keep_with_next=True
 return para
def h(t,level=1):return d.add_heading(t,level)
def link(label,url):
 p0=d.add_paragraph();a=OxmlElement('w:hyperlink');a.set(qn('r:id'),p0.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'185580');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);a.append(r);p0._p.append(a);return p0
def table(headers,rows,widths=None):
 tb=d.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
 widths=widths or [7.1/len(headers)]*len(headers)
 for col,w in zip(tb.columns,widths):col.width=Inches(w)
 pr=tb._tbl.tblPr;borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 pr.append(borders)
 for i,rowdata in enumerate([headers]+list(rows)):
  row=tb.rows[0] if i==0 else tb.add_row()
  trpr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
  if i==0:
   rep=OxmlElement('w:tblHeader');trpr.append(rep)
  for j,(cell,text) in enumerate(zip(row.cells,rowdata)):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for side,n in [('top',75),('bottom',75),('left',95),('right',95)]:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),str(n));e.set(qn('w:type'),'dxa');m.append(e)
   cp.append(m);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'173B57' if i==0 else ('F1F5F8' if i%2==0 else 'FFFFFF'));cp.append(shade)
   para=cell.paragraphs[0];para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.03
   run=para.add_run(str(text));run.font.size=Pt(10)
   if i==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
  if i==0:
   for cell in row.cells:cell.paragraphs[0].paragraph_format.keep_with_next=True
 p('') .paragraph_format.space_after=Pt(0)
 return tb

p('Данные автомобилей AutoHaina', 'Title')
p('Справочник для наполнения вашего сайта','Subtitle')
p('Проверка сайта 13–14 октября 2026 года. Итоговая версия от 14 октября 2026 года.')
p('Документ описывает данные, которые AutoHaina использует в каталоге и карточках автомобилей: поля объявления, фильтры, технические характеристики, оснащение, цены и справочники. Его можно передать разработчику как основу для формы добавления автомобиля, фильтрации и страницы подробностей.')
p('Собраны все 162 значения из публичного ответа справочника расширенных фильтров и состояний, а также 348 уникальных названий подробных параметров из девяти проверенных карточек. Приведён полный список 628 записей справочника марок, включая сомнительные служебные записи. Названия полей переведены и пояснены на русском; исходные подписи сохранены для сопоставления.')
p('Границы полноты. Это справочник структуры данных с примерами, а не выгрузка каждого объявления. Проверка девяти карточек не доказывает отсутствие дополнительных параметров в других комплектациях. Числа, свободные описания и перечни моделей не являются закрытыми справочниками. Закрытая административная база сайта не проверялась.')
h('Содержание',2)
for s in ['1 Поля объявления и краткой карточки','2 Фильтры и все полученные варианты значений','3 Подробные характеристики и примеры','4 Дополнительные данные карточки и расчёта цены','5 Ошибки источника и правила переноса','6 Проверенные автомобили и источники','7 Полный полученный справочник марок'] :p(s)
link('Основной источник — каталог AutoHaina','https://www.autohaina.com/vehicles')
h('1 Поля объявления и краткой карточки')
p('Ниже указаны фактические ключи публичного ответа карточки. Названия ключей принадлежат источнику. Типы хранения и пояснения — рекомендации для вашего сайта. Наличие поля у части машин не означает обязательность для всех автомобилей.')
toprows=[
('id','Идентификатор объявления','Строка; например 0md6fj8ctt. На странице выводится с префиксом AH.'),
('carCode','Код автомобиля у источника','Строка; может включать несколько кодов, разделённых знаком |.'),
('brandId','Идентификатор марки','Ссылка на справочник марок.'),
('title','Название объявления','Марка, серия, модельный год, двигатель и комплектация в свободном тексте.'),
('titleZh','Название на китайском','Исходное название; полезно для проверки перевода.'),
('metaTitle','Заголовок для поисковых систем','Отдельное текстовое поле.'),
('metaDescription','Описание для поисковых систем','Текст; может совпадать с замечаниями продавца.'),
('remark','Комментарий продавца','Состояние, ремонт, окрасы и другие особенности; свободный текст.'),
('featureTags','Метки автомобиля','Массив строк из справочника признаков.'),
('highlights','Краткие преимущества','Массив выделенных характеристик; может дублировать метки.'),
('highlightItems','Структурированные преимущества','Массив; в проверенных ответах встречался пустой.'),
('price','Цена автомобиля','Число в исходном ответе. В проверенных примерах базовые суммы соответствуют CNY; экран переводит их в выбранную валюту.'),
('fobPrice','Цена с обозначением FOB','Отдельное число в ответе. Не смешивать с ценой автомобиля и пользовательским расчётом.'),
('listedAt','Дата публикации','Дата объявления.'),
('registrationDate','Дата регистрации','Дата; также встречается внутри specs.'),
('horsepower','Мощность в лошадиных силах','Число; для гибрида уточнять, относится ли к ДВС или ко всей системе.'),
('year','Год','Строка с годом. Название модельного года в title может отличаться.'),
('mileage','Пробег','В ответе строка с единицей и разделителями, например 60,000 km; для поиска нужен отдельный числовой пробег в км.'),
('tag','Основная метка','Строка; может быть пустой.'),
('images','Фотогалерея','Массив URL изображений с порядком показа.'),
('specs','Краткие характеристики','Объект; полный набор обнаруженных вложенных ключей приведён далее.')]
table(['Ключ','Значение','Формат и пояснение'],toprows,[1.25,1.65,4.2])
specnames={'engine':'Двигатель','transmission':'Коробка передач','fuelType':'Топливо или тип силовой установки','driveType':'Привод','seats':'Количество мест','color':'Цвет кузова','bodyType':'Тип или класс кузова','displacement':'Объём и обозначение двигателя','flowMode':'Тип наддува','acceleration':'Разгон до 100 км/ч','fuelConsumption':'Расход топлива','batteryCapacity':'Ёмкость батареи','batteryType':'Химия батареи','range':'Запас хода','registrationDate':'Дата регистрации','officialPrice':'Цена производителя','condition':'Новый или подержанный','location':'Местонахождение'}
specrows=[]
for key,vals in U['specs'].items():
 ex=next((v for v in vals if v and str(v)!='0'),vals[0]);specrows.append((key,specnames[key],str(ex)))
h('Вложенные краткие характеристики',2);table(['Ключ в specs','Русское название','Пример из источника'],specrows,[1.5,3.0,2.6])
p('Отдельного закрытого списка объёмов двигателя в проверенном меню нет. Поле displacement содержит, например, 1.5, 1.5T, 2.0T, 4.0L и даже New energy у электромобиля. Подробные параметры различают объём в литрах и см³. Для вашего сайта следует хранить число и тип наддува отдельно.')
h('2 Фильтры и все полученные варианты значений')
p('Источник: каталог и публичный справочник /api/vehicles/filters?lang=en. Далее перечислены все 162 значения ответа, включая дубли и смешанные категории. Русское пояснение не заменяет исходное значение. Дополнительно в интерфейсе присутствуют интервалы цены, года, пробега и мощности.')
table(['Фильтр','Наблюдавшиеся варианты или поведение'],[
('Марка и серия','Марка из справочника; серия зависит от выбранной марки. Поиск по названию модели или ключевым словам.'),
('Состояние','Любое / новый / подержанный.'),
('Цена при валюте USD','Любая; менее $14.9K; $14.9K–$29.8K; $29.8K–$44.7K; $44.7K–$74.6K; $74.6K и выше; произвольный диапазон. Пороговые суммы зависят от конвертации.'),
('Год','Любой; 2026, 2025, 2024, 2023, 2022, 2021, 2020; произвольный диапазон. Это быстрые кнопки, а не все годы базы.'),
('Пробег','Любой; менее 10 000; 10 000–30 000; 30 000–60 000; 60 000–100 000; от 100 000 км; произвольный диапазон.'),
('Мощность','<160 HP; 160–200 HP; 200–300 HP; 300+ HP; произвольный диапазон. Точность включения границ в запрос не проверялась.'),
('Быстрые категории главной страницы','Новые; New Energy / EV; Fuel / Hybrid; SUV / Off-road / Pickup; более 160 HP; 160 HP и меньше; подержанные до 180 дней; подержанные 3–5 лет.')],[1.6,5.5])
fgroups={'conditions':'Состояние','fueltypes':'Топливо и силовая установка','gearboxes':'Коробки передач','driveTypes':'Привод','levels':'Кузов и размерный класс','bodyColors':'Цвет кузова','seatCounts':'Количество мест','batteryTypes':'Тип батареи','flowModes':'Наддув и впуск двигателя','featureTags':'Метки состояния и предложения'}
def ft(v):
 s=str(v)
 if s in F:return F[s]
 if isinstance(v,int):return str(v)+(' места' if v in [2,4] else ' мест')
 n=re.search(r'\d+',s);n=n.group() if n else ''
 if 'manual integration' in s:return n+' ступеней, автомат с ручным режимом'
 if 'DHT' in s:return n+' ступеней, гибридная трансмиссия DHT'
 if 'CVT' in s:
  nn=re.findall(r'\d+',s);return ('Гибридная e CVT' if 'E-CVT' in s else 'Вариатор CVT')+(', имитация '+nn[-1]+' передач' if nn else '')
 if 'wet dual-clutch' in s:return n+' ступеней, робот с мокрыми сцеплениями'
 if 'dry dual clutch' in s:return n+' ступеней, робот с сухими сцеплениями'
 if 'automatic' in s:return n+' ступеней, автомат'
 if 'manual' in s:return n+' ступеней, механическая коробка'
 return s
for k,vs in filters.items():
 h(fgroups[k],2)
 table(['Значение AutoHaina','Пояснение на русском'],[(v,ft(v)) for v in vs],[3.3,3.8])
p('В фильтре Body Type смешаны форма кузова, размерный класс и назначение. В вашем каталоге лучше разделить кузов, класс и категорию использования. Варианты батарей и коробок тоже содержат синонимы и разные уровни детализации.')
h('3 Подробные характеристики и примеры')
p('В девяти карточках обнаружено 37 названий разделов и 603 сочетания «раздел + название параметра». После удаления точных повторов названий осталось 348 уникальных подписей. Все они перечислены ниже один раз. Похожие подписи с разным написанием оставлены, чтобы не скрыть различия источника.')
p('В первой колонке дано русское значение и исходная подпись. Во второй — один пример из проверенной карточки и её код К1–К9; ссылки приведены в разделе 6. Примеры сохраняют исходные числа, единицы и формулировки. Они не являются перечнем всех допустимых значений. Параметры распределены по первому разделу, в котором встретились при сборе.')
p('Состояния оборудования в источнике: Standard equipment или ● — штатное оснащение; Optional equipment — опция; сочетания Front/Rear и Main/Deputy описывают разные места. Знак «-», отсутствие строки и ноль нельзя без проверки считать одним и тем же состоянием. В гарантиях и рекламных описаниях оставлены условия источника без подтверждения их применимости к экспорту.')
groupsru={'Basic Parameters':'Основные технические параметры','Car body':'Габариты и кузов','Engine':'Двигатель внутреннего сгорания','gearbox':'Коробка передач','Chassis steering':'Подвеска и рулевое управление','Wheel brake':'Колёса и тормоза','Electric motor':'Электропривод и параметры батареи','Active/Passive Safety Equipment':'Активная и пассивная безопасность','Auxiliary/Control Configuration':'Помощь при управлении и бездорожье','Exterior/Security Features':'Наружное оснащение и защита','Internal configuration':'Салон и внутреннее оснащение','Seat configuration':'Сиденья','Multimedia Configuration':'Мультимедиа','Lighting configuration':'Освещение','Glass/Rearview Mirror':'Стёкла и зеркала','Air conditioner/refrigerator':'Климат и холодильник','Battery/Charging':'Батарея и зарядка','Passive Safety':'Дополнительная пассивная безопасность','Active Safety':'Дополнительная активная безопасность','Driving Control':'Управление движением и подвеской','Driving hardware':'Датчики камеры и вычислители','Driving function':'Функции помощи водителю','Appearance/Anti-theft':'Внешнее оснащение и противоугонные функции','Exterior lighting':'Наружный свет','Sunroof/Glass':'Крыша и остекление','Outside rearview mirror':'Наружные зеркала','Interconnection/Intelligentization':'Связь и интеллектуальные системы','Steering wheel/Interior rearview mirror':'Руль и приборы','In-car charging':'Зарядка устройств в салоне','Audio/Interior Lighting':'Звук и салонное освещение','Optional package':'Фирменные дополнительные пакеты','Parameter Information':'Дополнительные технические поля','Security Configuration':'Дополнительные поля безопасности','Control Configuration':'Дополнительные поля управления','Intelligent/Assisted Driving':'Дополнительные ассистенты','External Configuration':'Дополнительное внешнее оборудование','Screen/System':'Дополнительные системные функции'}
seen=set();fieldcount=0
for g,items in U['groups'].items():
 rows=[]
 for key,values in items.items():
  if key in seen:continue
  seen.add(key);fieldcount+=1
  nonempty=[x for x in values if str(x['value']).strip() not in ['', '-', '0']]
  choice=min(nonempty or values,key=lambda x:len(str(x['value'])))
  val=str(choice['value']);source=codes[choice['source']]
  rows.append((T[key]+'\n'+key,val+'\nИсточник '+source))
 if rows:
  h(groupsru.get(g,g),2);p('Раздел источника: '+g+'.')
  table(['Характеристика и исходное название','Пример значения'],rows,[4.3,2.8])
assert fieldcount==348
h('4 Дополнительные данные карточки и расчёта цены')
table(['Блок','Данные'],[
('Навигация breadcrumb','brandId, brandName, seriesId, seriesName — идентификаторы и названия марки и серии.'),
('Подробные параметры parameters_by_group','Раздел → массив записей. Каждая запись содержит group_name, param_name, param_value, idx. idx задаёт порядок. Наборы разделов зависят от автомобиля.'),
('Метаданные geo_meta','В проверенных ответах: sourceName, cityZh, coverUrl, updatetime, verifiedAt, updatedAt, publishedAt. Это поля источника и обновлений; их смысл и достоверность не следует расширять до гарантии проверки машины.'),
('Менеджер sales_manager','Имя, отображаемое имя, изображение, описание, телефон, должность и ссылки мессенджеров. Для своего сайта нужны ваши контактные данные.'),
('Похожие автомобили relatedVehicles','Массив рекомендованных объявлений с их идентификаторами, названиями, ценами, фотографиями и частью характеристик.'),
('Условия предложения','Цена за единицу; минимальное количество MOQ; способы оплаты; режим «Inquiry» — цена по запросу; рекомендованная цена MSRP.'),
('Калькулятор экспорта','Режим EXW / FCA / FOB / CIF; цена машины, комиссия сервиса, внутренняя доставка, портовые расходы, обработка, международный фрахт, расчётная итоговая сумма.'),
('Валюта','Выбранная валюта, отображаемый обменный курс и пересчёт. На сайте указан источник open.er-api.com и кэш курса на один час.'),
('Действия посетителя','Избранное, сравнение, ссылка на объявление, обращение к менеджеру. Это состояние интерфейса или пользователя, а не технические характеристики машины.')],[2.05,5.05])
p('Цена Car Price Only относится к стоимости автомобиля. На странице прямо указано, что сервис, внутренняя логистика, портовые сборы, обработка и международная перевозка рассчитываются отдельно. Для цены по запросу нельзя автоматически показывать MSRP как цену продажи. Формулы включения расходов для каждого экспортного режима в этой проверке не тестировались.')
h('5 Ошибки источника и правила переноса')
issues=[
('Нулевые значения вместо отсутствующих','У eπ007 К2 краткое batteryCapacity равно 0, а подробное Battery energy равно 73.5 kWh. У Prado К9 seats равно 0. Ноль нельзя принимать за фактическую ёмкость или число мест.'),
('Противоречие кузова','В просмотренной карточке GLC К1 верхний блок BODY показывал Sedan, а specs.bodyType и подробный блок — SUV. Храните источник каждого значения и статус ручной проверки.'),
('Объём смешан с обозначением','2.0T объединяет литраж и наддув, а New energy вообще не является объёмом. Разделяйте объём в см³, объём в литрах, наддув и исходную строку.'),
('Год модели и регистрация','В названии GLC К1 есть 2022, поле year равно 2023, дата регистрации 2023-02-13. Эти значения описывают разные события; не объединяйте их в один год без правила.'),
('Неоднозначная мощность','У гибридов есть мощность ДВС, каждого электромотора и всей системы. Поле horsepower карточки не всегда означает суммарную мощность.'),
('Разные циклы измерения','CLTC, WLTC и NEDC для запаса хода и расхода нужно хранить отдельно вместе с единицей. Общий запас хода и электрический пробег — разные поля.'),
('Ошибки перевода','Precursor, increasing program, Non-operational, Business-to-non-profit и Rear seat adjustment method переведены неоднозначно. В частности Non-operational может означать некоммерческое использование, а не неисправность. Исходный китайский текст нужно проверять до публикации.'),
('Дубли в словарях','Plug-in hybrid и plug-in hybrid различаются только регистром. Automatic и конкретная 9-ступенчатая AT могут описывать один автомобиль на разных уровнях детализации.'),
('Справочник марок содержит мусор','Публичный ответ вернул 628 записей, включая пять Date Fallback Make без make_id. Наличие марки в справочнике не подтверждает наличие её автомобилей в продаже. Наблюдались также несоответствия подписей и изображений марок в меню.'),
('Изменяемые данные','Цены, курсы, наличие, даты, набор значений фильтров и объявления меняются. Сохраняйте дату получения и URL; не выдавайте снимок за постоянную базу.')]
for title,body in issues:
 pp=p('');rr=pp.add_run(title+'. ');rr.bold=True;pp.add_run(body)
h('Рекомендуемая структура для вашего сайта',2)
p('Это предложение по реализации, а не утверждение о внутренней базе AutoHaina. Отделите объявления о конкретных машинах от справочника моделей и комплектаций. У объявления храните цену, пробег, регистрацию, состояние, местонахождение, фотографии и замечания. У комплектации — заводские характеристики. Сохраняйте исходные значения рядом с нормализованными.')
table(['Сущность','Что хранить'],[
('Марка → серия → комплектация','Отдельные ID, названия, исходные названия, логотип и связи.'),
('Объявление','ID, связь с комплектацией, цена и валюта, статус цены, пробег, даты, состояние, местонахождение, описание.'),
('Характеристика','Стабильный собственный ключ, название на русском, группа, тип данных, единица измерения, применимость к ДВС / BEV / HEV / PHEV / EREV.'),
('Значение характеристики','Число или текст, исходная строка, цикл измерения при необходимости, URL источника, дата получения, статус проверки.'),
('Опция','Штатно / опционально / отсутствует / неизвестно; отдельно место установки, например водитель, пассажир, передний ряд, задний ряд.'),
('Справочник значений','Каноническое значение и все алиасы источника. Кузов, размерный класс и назначение — раздельные справочники.'),
('Медиа','URL или собственный файл, порядок, подпись, главное изображение. Права на повторную публикацию изображений в ходе этой проверки не устанавливались.')],[2.0,5.1])
h('6 Проверенные автомобили и источники')
p('Основные наблюдения сделаны в браузере 13 октября, публичные данные повторно получены 14 октября 2026 года. Коды К1–К9 служат ссылками для примеров в разделе 3. Покрыты бензиновый автомобиль, электромобили, подключаемые гибриды, автомобиль с увеличителем запаса хода, минивэн и внедорожники. Дизельные модели и пикапы отдельно не проверялись.')
for i,(f,slug,name) in enumerate(zip(files,slugs,names)):
 data=C[f]['data'];h(codes[f]+' '+name,2)
 p('ID '+data['id']+'; год в поле year: '+str(data['year'])+'; пробег: '+str(data['mileage'])+'.')
 p('Название источника: '+data['title'])
 link('Открыть карточку '+codes[f],'https://www.autohaina.com/vehicles/'+slug)
 link('Публичные данные карточки '+codes[f],'https://www.autohaina.com/api/vehicles/'+data['id']+'?lang=en')
h('Общие источники',2)
for label,url in [('Главная страница','https://www.autohaina.com/'),('Каталог и фильтры','https://www.autohaina.com/vehicles'),('Справочник значений фильтров','https://www.autohaina.com/api/vehicles/filters?lang=en'),('Русская версия справочника для сверки переводов','https://www.autohaina.com/api/vehicles/filters?lang=ru'),('Публичный справочник марок','https://www.autohaina.com/api/brands?lang=en')]:link(label,url)
h('7 Полный полученный справочник марок')
p('В ответе /api/brands?lang=en получено 628 записей. Поля: name — внутренний ID записи, make_name — название, make_logo — URL логотипа, make_id — внешний ID, firstletter — буква группировки, is_active — признак активности. Ниже приведены все 628 названий, отсортированные для чтения. HTML сущности декодированы, повторы названий сохранены. Это записи справочника, а не список марок с подтверждённым наличием машин.')
brands=load('brands.json')['data'];buckets=collections.defaultdict(list)
for b in brands:
 name=html.unescape(b['make_name']);letter=name[0].upper() if name else '?'
 if not ('A'<=letter<='Z'):letter='Прочие'
 buckets[letter].append(name+(' [нет make_id]' if not b.get('make_id') else ''))
for letter,values in sorted(buckets.items()):
 h(letter,2)
 # A compact dictionary is best read as prose separated by semicolons.
 vals=sorted(values,key=str.casefold)
 for start in range(0,len(vals),20):p('; '.join(vals[start:start+20])+'.')
p('Служебные записи Date Fallback Make не следует автоматически переносить в пользовательский выбор марок. Перечень серий для всех 628 записей в рамках этого документа не выгружался; он зависит от выбранной марки.')
for rootxml in [d.styles.element,d.element]:
 for border in list(rootxml.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.save(OUT)
print(OUT.resolve());print('Unique technical fields',fieldcount,'Filter entries',sum(map(len,filters.values())),'Brand rows',len(brands))
